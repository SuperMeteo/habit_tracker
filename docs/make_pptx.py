# -*- coding: utf-8 -*-
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "Leelawadee UI"
NAVY   = RGBColor(0x17, 0x37, 0x5E)
LIGHT  = RGBColor(0xF5, 0xF7, 0xFA)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREY   = RGBColor(0x5A, 0x6B, 0x7B)
DARK   = RGBColor(0x1B, 0x2B, 0x3A)
TEAL   = RGBColor(0x17, 0x8C, 0x82)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x1E, 0x88, 0x4B)
BLUE   = RGBColor(0x2E, 0x75, 0xB6)
AMBER  = RGBColor(0xC9, 0x82, 0x1A)
BORDER = RGBColor(0xD9, 0xDF, 0xE6)

# ---- auto-fit helper -------------------------------------------------------
import os as _os
from PIL import ImageFont as _IF
_REG = r"C:\Windows\Fonts\LeelawUI.ttf"
_BLD = r"C:\Windows\Fonts\LeelaUIb.ttf"
if not _os.path.exists(_BLD):
    _BLD = _REG
_fc = {}


def _font(sz, bold):
    k = (round(sz, 1), bold)
    if k not in _fc:
        _fc[k] = _IF.truetype(_BLD if bold else _REG, max(6, int(round(sz * 96.0 / 72.0))))
    return _fc[k]


def _nlines(text, f, max_px):
    if not text.strip():
        return 1
    lines, cur = 1, ""
    for w in text.split(" "):
        trial = w if not cur else cur + " " + w
        if f.getlength(trial) <= max_px:
            cur = trial
        else:
            while f.getlength(w) > max_px and len(w) > 1:
                cut = len(w)
                while cut > 1 and f.getlength(w[:cut]) > max_px:
                    cut -= 1
                lines += 1
                w = w[cut:]
            lines += 1
            cur = w
    return lines


def need_in(text, w_in, size, bold, space_after, line):
    """ความสูงที่ข้อความต้องใช้ (นิ้ว)"""
    f = _font(size, bold)
    max_px = w_in * 96.0 - 6
    total = 0.0
    for ln in text.split(chr(10)):
        n = _nlines(ln, f, max_px)
        total += n * size * 1.32 * line + space_after
    return total / 72.0


def fit_size(text, w_in, h_in, size, bold=False, space_after=4, line=1.0, floor=8.5):
    while size > floor and need_in(text, w_in, size, bold, space_after, line) > h_in:
        size -= 0.5
    return size


prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def txbox(slide, x, y, w, h, text, size=14, bold=False, color=DARK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4, line=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill=WHITE, line_col=BORDER, line_w=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_col is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_col
        s.line.width = Pt(line_w)
    if adj is not None and len(s.adjustments):
        s.adjustments[0] = adj
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, bar_color, title, body_lines, title_size=17, body_size=13):
    rect(slide, x, y, w, h, WHITE, BORDER, adj=0.04)
    rect(slide, x, y + 0.05, 0.075, h - 0.1, bar_color, None, shape=MSO_SHAPE.RECTANGLE)
    txbox(slide, x + 0.32, y + 0.22, w - 0.6, 0.4, title, title_size, True, DARK)
    if body_lines:
        bw, bh = w - 0.6, h - 1.02
        bs = fit_size(body_lines, bw, bh, body_size, False, 6, 1.15)
        txbox(slide, x + 0.32, y + 0.85, bw, bh, body_lines, bs, False,
              GREY, space_after=6, line=1.15)


def note(slide, x, y, w, h, label, text):
    rect(slide, x, y, w, h, RGBColor(0xFF, 0xFB, 0xEF), RGBColor(0xE8, 0xC4, 0x7A), adj=0.12)
    txbox(slide, x + 0.28, y + 0.16, 1.6, 0.3, label, 12, True, AMBER)
    bw, bh = w - 0.56, h - 0.68
    ns = fit_size(text, bw, bh, 13, False, 4, 1.15)
    txbox(slide, x + 0.28, y + 0.52, bw, bh, text, ns, False,
          RGBColor(0x6B, 0x54, 0x1E), line=1.15)


def note1(slide, x, y, w, label, text, h=0.55):
    """กล่องหมายเหตุแบบบรรทัดเดียว (ป้ายกำกับอยู่ซ้าย)"""
    rect(slide, x, y, w, h, RGBColor(0xFF, 0xFB, 0xEF), RGBColor(0xE8, 0xC4, 0x7A), adj=0.2)
    txbox(slide, x + 0.28, y + 0.17, 1.5, 0.3, label, 12, True, AMBER)
    bw = w - 2.1
    ts = fit_size(text, bw, 0.32, 12.5, False, 0, 1.0, floor=9)
    txbox(slide, x + 1.82, y + 0.18, bw, 0.32, text, ts, False, RGBColor(0x6B, 0x54, 0x1E))


def chrome(slide, title, subtitle, page):
    bg = rect(slide, 0, 0, 13.333, 7.5, LIGHT, None, shape=MSO_SHAPE.RECTANGLE)
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    top = rect(slide, 0, 0, 13.333, 0.17, NAVY, None, shape=MSO_SHAPE.RECTANGLE)
    slide.shapes._spTree.remove(top._element)
    slide.shapes._spTree.insert(3, top._element)
    txbox(slide, 0.55, 0.42, 11.5, 0.5, title, 27, True, NAVY)
    if subtitle:
        txbox(slide, 0.58, 1.03, 11.5, 0.3, subtitle, 12.5, False, GREY)
    txbox(slide, 0.55, 7.05, 6.0, 0.25, "\u0e42\u0e04\u0e23\u0e07\u0e01\u0e32\u0e23 1 | \u0e01\u0e32\u0e23\u0e19\u0e33\u0e40\u0e2a\u0e19\u0e2d\u0e1c\u0e25\u0e01\u0e32\u0e23\u0e14\u0e33\u0e40\u0e19\u0e34\u0e19\u0e07\u0e32\u0e19 \u2265 75%", 9.5, False, GREY)
    txbox(slide, 9.5, 7.05, 3.0, 0.25, "\u0e27\u0e34\u0e17\u0e22\u0e32\u0e25\u0e31\u0e22\u0e40\u0e17\u0e04\u0e19\u0e34\u0e04\u0e40\u0e0a\u0e35\u0e22\u0e07\u0e43\u0e2b\u0e21\u0e48", 9.5, False, GREY, align=PP_ALIGN.RIGHT)
    txbox(slide, 12.65, 7.03, 0.4, 0.25, str(page), 10, True, GREY, align=PP_ALIGN.RIGHT)


def new(title, subtitle, page):
    s = prs.slides.add_slide(BLANK)
    chrome(s, title, subtitle, page)
    return s


def table(slide, x, y, w, headers, rows, col_w, head_size=12.5, body_size=11.5, row_h=0.46):
    nrows = len(rows) + 1
    shp = slide.shapes.add_table(nrows, len(headers), Inches(x), Inches(y), Inches(w),
                                 Inches(0.5 + row_h * len(rows)))
    tbl = shp.table
    tbl.first_row = True
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(0.46)
    for i in range(1, nrows):
        tbl.rows[i].height = Inches(row_h)
    for c, htxt in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left = Inches(0.12)
        cell.margin_top = Inches(0.06)
        cell.margin_bottom = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = htxt
        r.font.name = FONT
        r.font.size = Pt(head_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xEE, 0xF2, 0xF7)
            cell.margin_left = Inches(0.12)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            txt, col, bold = (val if isinstance(val, tuple) else (val, DARK, False))
            p = cell.text_frame.paragraphs[0]
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = txt
            r.font.name = FONT
            r.font.size = Pt(body_size)
            r.font.bold = bold
            r.font.color.rgb = col
    return tbl


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, NAVY, None, shape=MSO_SHAPE.RECTANGLE)
tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(11.4), Inches(-0.4), Inches(2.6), Inches(2.2))
tri.fill.solid()
tri.fill.fore_color.rgb = RGBColor(0x1E, 0x4A, 0x77)
tri.line.fill.background()
tri.rotation = 90
tri.shadow.inherit = False
pie = s.shapes.add_shape(MSO_SHAPE.PIE, Inches(10.2), Inches(3.3), Inches(2.8), Inches(2.8))
pie.fill.solid()
pie.fill.fore_color.rgb = TEAL
pie.line.fill.background()
pie.shadow.inherit = False
try:
    pie.adjustments[0] = -90.0
    pie.adjustments[1] = -20.0
except Exception:
    pass

txbox(s, 0.75, 1.25, 9.5, 1.6,
      "แบบฟอร์มนำเสนอ\nผลการดำเนินงานโครงการ 1", 36, True, WHITE, line=1.25)
txbox(s, 0.78, 2.95, 9.0, 0.35,
      "สำหรับนักศึกษาที่มีความก้าวหน้าไม่น้อยกว่า 75%", 15, False, RGBColor(0xBF, 0xD3, 0xE6))
rect(s, 0.75, 3.55, 7.6, 1.55, None, RGBColor(0x4E, 0x74, 0x9B), 1.0, adj=0.08)
txbox(s, 1.05, 3.85, 7.1, 1.0,
      "ชื่อโครงงาน : Habit Tracker — แอปพลิเคชันติดตามพฤติกรรมประจำวันบนมือถือ\n"
      "ผู้จัดทำ : [ชื่อ–สกุล / รหัสนักศึกษา]\n"
      "อาจารย์ที่ปรึกษา : [ชื่ออาจารย์]",
      14, False, WHITE, space_after=7, line=1.2)
txbox(s, 0.78, 6.85, 10.0, 0.3,
      "การประเมินปลายภาค “โครงการ 1” | ผลงานสมบูรณ์ 100% และการสอบจบ นำเสนอในภาคเรียนถัดไป",
      11, False, RGBColor(0x9F, 0xB8, 0xD0))

# ---------------------------------------------------------------- 2. Problem
s = new("1. ปัญหา ที่มา และเป้าหมายของโครงงาน",
        "ทบทวนให้กรรมการเข้าใจว่า “ทำอะไร – เพื่อใคร – แก้ปัญหาอะไร” ภายใน 1 นาที", 2)
card(s, 0.55, 1.5, 3.95, 4.2, RED, "ปัญหา / ความจำเป็น",
     "• คนตั้งเป้าสร้างนิสัยใหม่ (ออกกำลังกาย อ่านหนังสือ ดื่มน้ำ) แล้วเลิกกลางคัน "
     "เพราะไม่มีเครื่องมือที่ทำให้เห็น “ความต่อเนื่อง” ของตัวเอง\n"
     "• จดในสมุดหรือโน้ตในมือถือ สรุปย้อนหลังไม่ได้ ไม่รู้ว่าทำได้กี่วันติดกัน\n"
     "• แอปในตลาดส่วนใหญ่ต้องสมัครสมาชิก ต่ออินเทอร์เน็ต และเป็นภาษาอังกฤษ\n\n"
     "กลุ่มเป้าหมาย : นักเรียน–นักศึกษา และวัยทำงานที่ต้องการสร้างวินัยส่วนตัว")
card(s, 4.72, 1.5, 3.95, 4.2, TEAL, "วัตถุประสงค์",
     "1) พัฒนาแอปบนมือถือด้วย Flutter สำหรับบันทึกพฤติกรรมประจำวัน "
     "ที่ทำงานแบบออฟไลน์ ไม่ต้องสมัครสมาชิก\n\n"
     "2) คำนวณสถิติจากข้อมูลที่บันทึก ได้แก่ จำนวนวันต่อเนื่อง (Streak) "
     "และอัตราความสำเร็จรายสัปดาห์ พร้อมแสดงผลเป็นกราฟและปฏิทินสี\n\n"
     "3) รองรับพฤติกรรมทั้งแบบ “ทำ / ไม่ทำ” และแบบเชิงตัวเลขที่มีเป้าหมาย "
     "เช่น ดื่มน้ำ 8 แก้ว วิ่ง 5 กม. พร้อมแม่แบบสำเร็จรูปให้เลือก")
card(s, 8.89, 1.5, 3.89, 4.2, BLUE, "ขอบเขตงาน",
     "อยู่ในขอบเขต\n"
     "• แอปบน Android (โค้ดชุดเดียว ต่อยอด iOS ได้)\n"
     "• เก็บข้อมูลในเครื่องด้วย SQLite 3 ตาราง\n"
     "• 3 หน้าจอหลัก : หน้าหลัก / สถิติ / ตั้งค่า\n"
     "• หน้าจอภาษาไทย + ธีมสว่างและมืด\n\n"
     "อยู่นอกขอบเขต\n"
     "• ระบบสมาชิก การซิงก์ขึ้นคลาวด์ และระบบแชร์กับเพื่อน")
note(s, 0.55, 5.9, 12.23, 1.05, "สิ่งที่ต้องใส่",
     "ใช้ข้อความสั้น กระชับ ไม่คัดลอกบทที่ 1 มาวางทั้งหน้า และควรมีภาพปัญหาหรือภาพบริบทจริง 1 ภาพหากมี\n"
     "แนะนำ : ใส่ภาพสมุดจดนิสัยหรือโน้ตในมือถือ เทียบกับภาพหน้าจอแอปที่ทำเสร็จแล้ว 1 ภาพ")

# ---------------------------------------------------------------- 3. Progress
s = new("2. สรุปความก้าวหน้าปัจจุบัน",
        "แสดงสถานะจริง ณ วันสอบ ไม่จำเป็นต้องครบ 100% แต่ต้องไม่น้อยกว่า 75% และมีหลักฐานตรวจสอบ", 3)
txbox(s, 0.55, 1.45, 4.0, 0.35, "ความก้าวหน้ารวม", 17, True, DARK)
rect(s, 0.55, 1.9, 12.23, 0.55, RGBColor(0xE3, 0xE9, 0xF0), None, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0.55, 1.9, 9.17, 0.55, TEAL, None, shape=MSO_SHAPE.RECTANGLE)
txbox(s, 10.0, 1.97, 2.0, 0.4, "75%", 20, True, NAVY)
table(s, 0.55, 2.58, 12.23,
      ["งาน / โมดูล", "สถานะ", "หลักฐาน"],
      [["โครงสร้างโปรเจกต์ + ออกแบบฐานข้อมูล 3 ตาราง (Categories / Habits / HabitLogs)",
        ("เสร็จแล้ว", GREEN, True), "โค้ด drift + ไฟล์ app_database.g.dart ที่สร้างอัตโนมัติ"],
       ["หน้าหลัก — เลือกวัน กดเช็ค habit กรอกค่าตัวเลข และแถบสรุปความคืบหน้าของวัน",
        ("เสร็จแล้ว", GREEN, True), "สาธิตบนอีมูเลเตอร์ / dashboard_screen.dart (283 บรรทัด)"],
       ["เพิ่ม–แก้ไข habit + แม่แบบสำเร็จรูป 18 รายการ 3 หมวด (กีฬา / ชีวิต / การศึกษา)",
        ("เสร็จแล้ว", GREEN, True), "add_habit_screen.dart + template_picker_screen.dart"],
       ["หน้าสถิติ — Streak, ปฏิทินความถี่ย้อนหลัง 1 ปี, กราฟแท่งรายสัปดาห์",
        ("เสร็จแล้ว", GREEN, True), "analytics_screen.dart + fl_chart / ภาพหน้าจอ"],
       ["ตั้งค่า — ธีมสว่าง / มืด / ตามระบบ และจำค่าไว้หลังปิดแอป",
        ("เสร็จแล้ว", GREEN, True), "settings_screen.dart + shared_preferences"],
       ["ระบบแจ้งเตือนตามเวลาที่ตั้งไว้",
        ("กำลังดำเนินการ", AMBER, True), "เก็บเวลาเตือนลงฐานข้อมูลได้แล้ว แต่ยังไม่ยิงแจ้งเตือนจริง"],
       ["ความถี่แบบ “N ครั้งต่อสัปดาห์”",
        ("กำลังดำเนินการ", AMBER, True), "มีตัวเลือกในหน้าเพิ่ม habit แล้ว แต่ยังนับผลเหมือนรายวัน"],
       ["ชุดทดสอบอัตโนมัติ (Unit / Widget test)",
        ("ยังไม่เสร็จ", RED, True), "test/widget_test.dart ยังเป็นไฟล์ตัวอย่างที่ Flutter สร้างให้"]],
      [6.1, 1.85, 4.28], row_h=0.40, body_size=10.5)
note1(s, 0.55, 6.4, 12.23, "สิ่งที่ต้องใส่",
      "75% : 5 โมดูลหลักเสร็จและสาธิตได้ · 2 รายการทำไปบางส่วน · ชุดทดสอบอัตโนมัติยังไม่เริ่ม")

# ---------------------------------------------------------------- 4. Design
s = new("3. การออกแบบและวิธีดำเนินงาน", "สถาปัตยกรรมแอปและฐานข้อมูลที่ใช้จริงในโครงงาน", 4)
rect(s, 0.55, 1.5, 6.6, 5.2, WHITE, BORDER, adj=0.03)
txbox(s, 0.85, 1.72, 6.0, 0.3, "สถาปัตยกรรมระบบ (Layered Architecture)", 15, True, DARK)


def layer(x, y, w, h, col, title, sub):
    rect(s, x, y, w, h, col, None, adj=0.14)
    txbox(s, x + 0.12, y + 0.13, w - 0.24, 0.28, title, 12, True, WHITE, align=PP_ALIGN.CENTER)
    if sub:
        ss = fit_size(sub, w - 0.28, 0.36, 9, False, 4, 1.05, floor=7)
        txbox(s, x + 0.14, y + 0.42, w - 0.28, 0.36, sub, ss, False, RGBColor(0xE4, 0xEE, 0xF7),
              align=PP_ALIGN.CENTER, line=1.05)


def arrow(x, y):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(0.3), Inches(0.22))
    a.fill.solid()
    a.fill.fore_color.rgb = RGBColor(0xA8, 0xB6, 0xC4)
    a.line.fill.background()
    a.shadow.inherit = False


layer(0.85, 2.15, 6.0, 0.85, BLUE, "ชั้นแสดงผล (UI) — Flutter Widgets",
      "หน้าหลัก · สถิติ · เพิ่ม/แก้ไข habit · เลือกแม่แบบ · ตั้งค่า   |   นำทางด้วย go_router (ShellRoute + แถบล่าง)")
arrow(3.85, 3.06)
layer(0.85, 3.35, 6.0, 0.85, TEAL, "ชั้นสถานะและตรรกะ — Riverpod Providers",
      "habitsProvider · logsForDateProvider · dashboardProvider · analyticsProvider · heatmapProvider   |   StreakCalculator")
arrow(3.85, 4.26)
layer(0.85, 4.55, 6.0, 0.85, RGBColor(0x6D, 0x4A, 0xAF), "ชั้นข้อมูล — AppDatabase (drift ORM)",
      "watchActiveHabits() · watchLogsForDate() · upsertLog() · getAllLogsForHabit()")
arrow(3.85, 5.46)
layer(0.85, 5.75, 6.0, 0.75, RGBColor(0x3F, 0x4B, 0x58), "ฐานข้อมูล SQLite ในเครื่อง — habit_tracker.db",
      "Categories (1) —< Habits (1) —< HabitLogs   |   คีย์ไม่ซ้ำ : (habitId, loggedDate)")

card(s, 7.35, 1.5, 5.43, 5.2, TEAL, "แนวทางและเครื่องมือที่ใช้",
     "เครื่องมือ / ภาษา / Framework\n"
     "• Flutter 3.44.8 + Dart — เขียนโค้ดชุดเดียวใช้ได้หลายแพลตฟอร์ม\n"
     "• drift 2.21 + sqlite3_flutter_libs — ฐานข้อมูลในเครื่อง เขียนคำสั่งค้นข้อมูลแบบตรวจชนิดข้อมูลให้ตั้งแต่ตอนคอมไพล์\n"
     "• flutter_riverpod 2.6 — จัดการสถานะและรีเฟรชหน้าจออัตโนมัติ\n"
     "• go_router 14.6 — จัดการเส้นทางหน้าจอ · fl_chart 0.69 — กราฟแท่ง\n"
     "• intl — วันที่ภาษาไทย · shared_preferences — จำค่าธีม\n\n"
     "หลักการทำงานของส่วนสำคัญ\n"
     "• หน้าจอ “เฝ้าดู” ฐานข้อมูลผ่าน Stream เมื่อกดเช็ค habit ข้อมูลจะถูกเขียนลง SQLite แล้วทุกหน้าที่เกี่ยวข้องอัปเดตเองทันที\n"
     "• ตัวคำนวณ Streak ไล่ย้อนหลังทีละวันสูงสุด 365 วัน ข้ามวันที่ไม่ใช่วันเป้าหมาย และไม่นับว่าขาดถ้าวันนี้ยังไม่ได้บันทึก\n\n"
     "เหตุผลที่เลือกแนวทางนี้\n"
     "• เก็บข้อมูลในเครื่อง = ใช้งานได้แม้ไม่มีอินเทอร์เน็ต ไม่ต้องมีเซิร์ฟเวอร์ และข้อมูลส่วนตัวไม่ออกจากเครื่องผู้ใช้",
     body_size=11)

# ---------------------------------------------------------------- 5. Results
s = new("4. ผลงานที่ดำเนินการเสร็จแล้ว", "เน้น “ของจริงที่ทำได้แล้ว” มากกว่าคำอธิบาย", 5)


SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "screenshots")
EVID = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_evidence")
T4_RESULT = "Streak ปัจจุบัน 3 วัน (หลังแก้บั๊กตัวคำนวณ)"
T4_STATUS = ("ผ่าน", GREEN, True)


def result_card(x, title, imgs, desc):
    card(s, x, 1.5, 3.95, 4.7, GREEN, title, "")
    ih = 2.75
    iw = ih * 824 / 1760
    gap = 0.2
    total = iw * len(imgs) + gap * (len(imgs) - 1)
    ix = x + 0.3 + (3.35 - total) / 2
    for name in imgs:
        rect(s, ix - 0.03, 2.02, iw + 0.06, ih + 0.06, WHITE, RGBColor(0xC6, 0xD1, 0xDC), adj=0.04)
        s.shapes.add_picture(os.path.join(SHOTS, name), Inches(ix), Inches(2.05), Inches(iw), Inches(ih))
        ix += iw + gap
    ds = fit_size(desc, 3.35, 1.3, 11.5, False, 4, 1.15)
    txbox(s, x + 0.32, 4.92, 3.35, 1.3, desc, ds, False, GREY, line=1.15)


result_card(0.55, "ผลงานส่วนที่ 1 — หน้าหลัก",
            ["01_dashboard.png", "05_settings.png"],
            "เลือกวันจากแถบวันที่ · กดเช็ค habit ของวันนั้น · habit เชิงตัวเลขกรอกค่าได้ "
            "(เช่น ดื่มน้ำ 6/8 แก้ว) · แถบสรุปความคืบหน้าของวัน\nใช้งานได้เต็มรูปแบบ บันทึกลงฐานข้อมูลจริง")
result_card(4.72, "ผลงานส่วนที่ 2 — จัดการ Habit",
            ["02_add_habit.png", "02b_templates.png"],
            "แม่แบบสำเร็จรูป 18 รายการ 3 หมวด · ตั้งชื่อ ไอคอน สี หมวดหมู่ · "
            "เลือกความถี่รายวันหรือเจาะจงวัน · ตั้งเป้าเชิงตัวเลขและหน่วย · แก้ไขและลบ habit เดิมได้")
result_card(8.89, "ผลงานส่วนที่ 3 — สถิติ",
            ["03_analytics.png", "03b_analytics_scroll.png"],
            "การ์ด Streak ปัจจุบันและสถิติสูงสุด · อัตราสำเร็จรายสัปดาห์ · "
            "ปฏิทินความถี่ย้อนหลัง 1 ปี · กราฟแท่งรายสัปดาห์ของ habit เชิงตัวเลข")
note1(s, 0.55, 6.3, 12.23, "สิ่งที่ต้องใส่",
      "หลักฐาน : ซอร์สโค้ดบน Git · ภาพหน้าจอจากแอปที่รันจริง (ข้อมูลตัวอย่าง) · สาธิตสดได้ในวันสอบ")

# ---------------------------------------------------------------- 6. Testing
s = new("5. การทดสอบและผลที่ได้", "แสดงหลักฐานว่าผลงานไม่ได้เพียง “ทำเสร็จ” แต่ผ่านการทดสอบบางส่วนแล้ว", 6)
table(s, 0.55, 1.45, 12.23,
      ["รายการทดสอบ", "วิธีทดสอบ", "ผลที่คาดหวัง", "ผลที่ได้", "สถานะ"],
      [["T1 ติดตั้งและรันแอป", "สั่ง flutter run บนอีมูเลเตอร์ Pixel 6",
        "แอปเปิดขึ้นหน้าหลักโดยไม่ค้าง", "คอมไพล์ผ่านและเปิดใช้งานได้", ("ผ่าน", GREEN, True)],
       ["T2 สร้างฐานข้อมูลครั้งแรก", "ลบแอปแล้วติดตั้งใหม่ เปิดหน้าเพิ่ม habit",
        "มีหมวดหมู่ตั้งต้น 5 หมวด", "ขึ้นครบ 5 หมวด", ("ผ่าน", GREEN, True)],
       ["T3 บันทึก habit รายวัน", "กดเช็ค habit แล้วปิดแอปและเปิดใหม่",
        "สถานะที่เช็คไว้ยังอยู่ครบ", "2/5 → 3/5 และหลังเปิดใหม่ยังเป็น 3/5", ("ผ่าน", GREEN, True)],
       ["T4 คำนวณ Streak", "เช็คติดต่อกัน 3 วัน แล้วเปิดหน้าสถิติ",
        "Streak ปัจจุบันแสดงค่า 3", T4_RESULT, T4_STATUS],
       ["T5 habit เชิงตัวเลข", "กรอกค่า 7 จากเป้า 8 แก้ว แล้วดูกราฟรายสัปดาห์",
        "หน้าหลักแสดง 7/8 และกราฟวันนั้นที่ค่า 7", "แสดง 7/8 แก้ว กราฟขึ้น 7 แก้ว", ("ผ่าน", GREEN, True)],
       ["T6 สลับธีมสว่าง / มืด", "ตั้งค่า เลือกธีมมืด ปิดแอปแล้วเปิดใหม่",
        "แอปยังคงเป็นธีมมืด", "ยังเป็นธีมมืดหลังเปิดใหม่", ("ผ่าน", GREEN, True)]],
      [2.35, 3.3, 2.75, 2.33, 1.5], row_h=0.5, body_size=10.5)
card(s, 0.55, 4.98, 6.1, 1.68, BLUE, "หลักฐาน", "", title_size=15)
_eh = 1.2
_ew = _eh * 824 / 1760
for _i, (_img, _lab) in enumerate([("t3_after_reload.png", "T3"), ("t4_stats.png", "T4"),
                                   ("t5_chart.png", "T5"), ("t6_after_reload.png", "T6")]):
    _x = 2.2 + _i * 1.1
    s.shapes.add_picture(os.path.join(EVID, _img), Inches(_x), Inches(5.3), Inches(_ew), Inches(_eh))
    txbox(s, _x + _ew + 0.05, 5.35, 0.45, 0.25, _lab, 10, True, BLUE)
card(s, 6.68, 4.98, 6.1, 1.68, AMBER, "ข้อค้นพบ",
     "• ทำงานได้ดี : อ่าน–เขียนข้อมูลกับ SQLite เร็ว หน้าจออัปเดตทันทีที่กดเช็ค\n"
     "• ต้องปรับปรุงก่อนสอบจบ : ระบบแจ้งเตือน ความถี่ N ครั้งต่อสัปดาห์ และชุดทดสอบอัตโนมัติ",
     title_size=15, body_size=12)
txbox(s, 0.55, 6.75, 12.23, 0.3,
      "วิธีทดสอบ T3–T6 : สั่งเบราว์เซอร์กดแอปจริงแบบอัตโนมัติ (ข้อมูลตัวอย่าง) แล้วเก็บภาพหน้าจอทุกขั้น",
      10.5, False, GREY)

# ---------------------------------------------------------------- 7. Problems
s = new("6. ปัญหา อุปสรรค และการแก้ไข", "กรรมการต้องเห็นกระบวนการคิดและการแก้ปัญหาระหว่างพัฒนา", 7)
table(s, 0.55, 1.5, 12.23,
      ["ปัญหาที่พบ", "สาเหตุ", "แนวทางแก้ไขที่ดำเนินการ", "ผลหลังแก้ไข"],
      [["ติดตั้งไลบรารีไม่ผ่าน ขึ้นข้อความ requires symlink support",
        "Windows ยังไม่ได้เปิดโหมดนักพัฒนา จึงสร้างลิงก์ไฟล์ที่ Flutter ต้องใช้ไม่ได้",
        "เปิด Developer Mode ในตั้งค่าของ Windows แล้วสั่งติดตั้งไลบรารีใหม่",
        "ติดตั้งครบและคอมไพล์ต่อได้"],
       ["ตรวจเครื่องมือด้วย flutter doctor ไม่ผ่านหัวข้อ Android",
        "ติดตั้ง Android Studio ไว้คนละตำแหน่งกับค่ามาตรฐาน และตัวติดตั้งไม่ได้แถมชุดคำสั่ง cmdline-tools มาด้วย",
        "ตั้งค่า JAVA_HOME ให้ชี้ไปที่โฟลเดอร์ jbr ของ Android Studio และดาวน์โหลด cmdline-tools แยกมาแตกไฟล์ลงในโฟลเดอร์ SDK",
        "flutter doctor ผ่าน และยอมรับ license ได้"],
       ["อีมูเลเตอร์ดับเองระหว่างคอมไพล์ครั้งแรก ซึ่งใช้เวลาราว 17 นาที",
        "หน่วยความจำเครื่องเหลือน้อย อีมูเลเตอร์กับตัวคอมไพล์ Gradle แย่งทรัพยากรกัน",
        "ปิดโปรแกรมอื่นก่อนคอมไพล์ และรอให้คอมไพล์เสร็จก่อนจึงเปิดอีมูเลเตอร์",
        "คอมไพล์รอบถัดไปใช้เวลาไม่ถึง 1 นาที"],
       ["คอมไพล์ไม่ผ่าน หาคลาส Habit และ HabitLog ไม่เจอ",
        "drift ต้องสร้างไฟล์โค้ดอัตโนมัติ (.g.dart) ก่อน จึงจะมีคลาสเหล่านี้ให้เรียกใช้",
        "สั่ง build_runner สร้างโค้ดใหม่ทุกครั้งที่แก้โครงสร้างตาราง และบันทึกขั้นตอนนี้ไว้ในคู่มือติดตั้ง",
        "คอมไพล์ผ่าน และเพื่อนที่รับโปรเจกต์ไปทำตามได้"],
       ["จำนวนวันต่อเนื่อง (Streak) ถูกรีเซ็ตเป็น 0 ทั้งที่ยังไม่ได้ขาด",
        "ตัวคำนวณนับวันปัจจุบันที่ยังไม่ได้กดเช็คเป็นวันที่ขาดทันที",
        "แก้ตรรกะให้ข้ามวันปัจจุบันถ้ายังไม่ได้บันทึก แล้วเริ่มนับต่อเนื่องจากเมื่อวานแทน",
        "Streak แสดงค่าถูกต้องตลอดทั้งวัน"]],
      [2.6, 3.1, 3.85, 2.68], row_h=0.83, body_size=10.5)
note1(s, 0.55, 6.32, 12.23, "สิ่งที่ต้องใส่",
      "เลือกเฉพาะประเด็นสำคัญ และอธิบายให้ครบทั้ง 4 ช่อง : อาการเป็นอย่างไร เกิดจากอะไร แก้อย่างไร ผลหลังแก้เป็นอย่างไร")
# ---------------------------------------------------------------- 8. Remaining
s = new("7. งานที่ยังเหลือและแผนพัฒนาสู่ 100%",
        "ส่วนนี้เชื่อมจาก “โครงการ 1” ไปสู่การทำผลงานสมบูรณ์และสอบจบในภาคเรียนหน้า", 8)
table(s, 0.55, 1.5, 12.23,
      ["งานที่เหลือ", "สถานะปัจจุบัน", "สิ่งที่จะดำเนินการต่อ", "เป้าหมาย / หลักฐานเมื่อเสร็จ"],
      [["ระบบแจ้งเตือนตามเวลาที่ผู้ใช้ตั้ง",
        ("20%", AMBER, True),
        "เขียนโค้ดตั้งเวลาแจ้งเตือนซ้ำทุกวันตามเวลาที่บันทึกไว้ และขอสิทธิ์แจ้งเตือนบน Android 13 ขึ้นไป",
        "สาธิตการแจ้งเตือนเด้งขึ้นจริงบนอีมูเลเตอร์"],
       ["ความถี่แบบ “N ครั้งต่อสัปดาห์”",
        ("40%", AMBER, True),
        "แก้ตัวคำนวณให้นับจำนวนครั้งที่ทำได้ในสัปดาห์เทียบกับเป้าหมาย แทนการนับรายวัน ทั้งหน้าหลักและหน้าสถิติ",
        "ผลทดสอบ : ตั้งเป้า 3 ครั้ง/สัปดาห์ ทำได้ 2 ครั้ง ต้องแสดง 2/3"],
       ["หน้าจัดการหมวดหมู่",
        ("ยังไม่เริ่ม", RED, True),
        "เพิ่มหน้าสำหรับเพิ่ม แก้ไข และลบหมวดหมู่ ปัจจุบันใช้ได้เฉพาะ 5 หมวดตั้งต้นที่ระบบสร้างให้",
        "ภาพหน้าจอหน้าจัดการหมวดหมู่ที่ใช้งานได้จริง"],
       ["ชุดทดสอบอัตโนมัติ",
        ("ยังไม่เริ่ม", RED, True),
        "เขียนชุดทดสอบตัวคำนวณ Streak (ทำต่อเนื่อง / ขาดกลางทาง / เจาะจงวัน) และทดสอบหน้าหลัก แทนไฟล์ตัวอย่างเดิม",
        "ผลรัน flutter test ผ่านทุกกรณี พร้อมภาพผลลัพธ์"],
       ["ส่งออกข้อมูลและการเผยแพร่แอป",
        ("ยังไม่เริ่ม", RED, True),
        "เพิ่มการส่งออกข้อมูลเป็นไฟล์ CSV/JSON ตั้งชื่อและไอคอนแอป แล้วสร้างไฟล์ APK สำหรับติดตั้งจริง",
        "ไฟล์ APK ที่ติดตั้งบนมือถือจริงได้ และไฟล์ข้อมูลที่ส่งออก"]],
      [2.5, 1.35, 5.2, 3.18], row_h=0.83, body_size=10.5)
note1(s, 0.55, 6.3, 12.23, "สิ่งที่ต้องใส่",
      "ระบุงานที่เหลือเป็นรูปธรรม : จะเพิ่มอะไร แก้ตรงไหน ทดสอบด้วยวิธีใด ไม่ใช้ถ้อยคำกว้าง ๆ เช่น “ปรับปรุงระบบให้สมบูรณ์”")
# ---------------------------------------------------------------- 9. Roles
s = new("8. การแบ่งหน้าที่และผลงานรายบุคคล",
        "ใช้ตรวจสอบการมีส่วนร่วมของสมาชิกแต่ละคนในการตัดเกรดโครงการ 1", 9)
table(s, 0.55, 1.5, 12.23,
      ["ชื่อสมาชิก", "หน้าที่รับผิดชอบ", "สิ่งที่ทำสำเร็จ", "หลักฐาน"],
      [["[ชื่อคนที่ 1]", "ออกแบบฐานข้อมูลและชั้นข้อมูล",
        "ตาราง Categories / Habits / HabitLogs · คำสั่งอ่าน–เขียนข้อมูล · ตัวคำนวณ Streak",
        "[Git commit / app_database.dart / streak_calculator.dart]"],
       ["[ชื่อคนที่ 2]", "ออกแบบและพัฒนาส่วนติดต่อผู้ใช้",
        "หน้าหลัก · หน้าเพิ่มและแก้ไข habit · หน้าเลือกแม่แบบ · ธีมสว่างและมืด",
        "[Git commit / ภาพหน้าจอ / สาธิตสด]"],
       ["[ชื่อคนที่ 3 ถ้ามี]", "หน้าสถิติ การทดสอบ และเอกสาร",
        "ปฏิทินความถี่ · กราฟรายสัปดาห์ · การ์ด Streak · คู่มือติดตั้ง SETUP.md",
        "[Git commit / ผลทดสอบ / เอกสาร]"]],
      [2.0, 3.0, 4.3, 2.93], row_h=0.85, body_size=11)
note(s, 0.55, 4.9, 12.23, 1.05, "สิ่งที่ต้องใส่",
     "สมาชิกทุกคนต้องตอบคำถามเกี่ยวกับงานของตนเองได้ และกรรมการอาจให้คะแนนรายบุคคลแตกต่างกันตามหลักฐานและความเข้าใจ\n"
     "ให้ปรับชื่อและหน้าที่ให้ตรงกับความจริง พร้อมเตรียมอธิบายโค้ดส่วนที่ตนเองเขียนได้")
rect(s, 0.55, 5.98, 12.23, 1.05, WHITE, BORDER, adj=0.1)
rect(s, 0.55, 6.03, 0.075, 0.95, TEAL, None, shape=MSO_SHAPE.RECTANGLE)
txbox(s, 0.87, 6.16, 3.0, 0.3, "คำถามที่ทุกคนควรตอบได้", 14, True, DARK)
txbox(s, 0.87, 6.56, 11.6, 0.36,
      "ข้อมูลถูกเก็บไว้ที่ไหน · Streak คำนวณอย่างไร · ทำไมเลือกเก็บข้อมูลในเครื่องแทนคลาวด์ · หน้าจออัปเดตเองได้อย่างไรเมื่อกดเช็ค",
      12, False, GREY)

# ---------------------------------------------------------------- 10. Summary
s = new("9. สรุปผลการดำเนินงานและสาธิต", "ปิดการนำเสนอด้วยสิ่งที่ทำได้จริง ณ วันสอบ", 10)
card(s, 0.55, 1.5, 3.95, 4.15, TEAL, "ระดับความสำเร็จ",
     "ความก้าวหน้าปัจจุบัน : 75%\n\n"
     "ส่วนสำคัญที่ทำสำเร็จ\n"
     "1) บันทึกพฤติกรรมประจำวันลงฐานข้อมูลในเครื่องได้จริง ทั้งแบบทำ/ไม่ทำ และแบบมีเป้าตัวเลข\n"
     "2) คำนวณและแสดง Streak อัตราสำเร็จรายสัปดาห์ และปฏิทินความถี่ย้อนหลัง 1 ปี\n"
     "3) มีแม่แบบ habit 18 รายการ และธีมสว่าง–มืดที่จำค่าไว้ได้")
card(s, 4.72, 1.5, 3.95, 4.15, GREEN, "สิ่งที่พิสูจน์ได้",
     "• แอปทำงานได้จริงบน Android Emulator (Pixel 6) — สาธิตสดได้\n"
     "• ฐานข้อมูล SQLite 3 ตาราง พร้อมหมวดหมู่ตั้งต้น 5 หมวด\n"
     "• ซอร์สโค้ด 26 ไฟล์ Dart ประมาณ 3,000 บรรทัด อยู่บน Git\n"
     "• คู่มือติดตั้ง SETUP.md ที่ผู้อื่นทำตามแล้วรันได้\n\n"
     "อ้างอิงหลักฐานในวันสอบ : ภาพหน้าจอ · การสาธิตสด · ประวัติการ commit")
card(s, 8.89, 1.5, 3.89, 4.15, BLUE, "เป้าหมายภาคเรียนหน้า",
     "• ทำระบบแจ้งเตือนและความถี่ N ครั้งต่อสัปดาห์ให้ครบ\n"
     "• เพิ่มหน้าจัดการหมวดหมู่ และการส่งออก/สำรองข้อมูล\n"
     "• เขียนชุดทดสอบอัตโนมัติและรันผ่านทั้งหมด\n"
     "• ทดลองใช้กับผู้ใช้จริงและประเมินความพึงพอใจ\n"
     "• สร้างไฟล์ APK ติดตั้งบนมือถือจริง และจัดทำรายงานฉบับสมบูรณ์")
rect(s, 0.55, 5.7, 12.23, 0.88, RGBColor(0xE8, 0xF5, 0xF3), RGBColor(0x9F, 0xD3, 0xCC), adj=0.2)
txbox(s, 0.95, 5.97, 2.7, 0.4, "DEMO ผลงานจริง", 19, True, TEAL)
txbox(s, 3.7, 6.03, 8.8, 0.4,
      "สาธิต 2–3 นาที : เพิ่ม habit จากแม่แบบ → กดเช็คของวันนี้ → เปิดหน้าสถิติดู Streak และปฏิทิน  |  เตรียมวิดีโอสำรองไว้ด้วย",
      12.5, False, RGBColor(0x2C, 0x5F, 0x5A))
txbox(s, 0.55, 6.6, 12.23, 0.4, "ขอบคุณครับ / ค่ะ   |   คำถามจากกรรมการ", 18, True, NAVY,
      align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 11. Tips
s = new("คำแนะนำก่อนวันสอบ", "สไลด์นี้ใช้เป็นแนวทางเตรียมตัว สามารถลบออกก่อนนำเสนอได้", 11)
txbox(s, 1.1, 1.85, 11.2, 3.3,
      "•  เวลานำเสนอแนะนำ 7–10 นาที และตอบคำถามประมาณ 3–5 นาที\n"
      "•  นำเสนอเฉพาะผลการดำเนินงานที่ทำได้จริง ณ วันสอบ ไม่จำเป็นต้องอ้างว่าโครงงานสมบูรณ์ 100%\n"
      "•  หลักฐานที่ต้องเตรียม : ซอร์สโค้ดบน Git · ภาพหน้าจอทุกหน้า · log จากการรันแอป · ไฟล์ SETUP.md · วิดีโอสาธิต\n"
      "•  เปิดอีมูเลเตอร์ทิ้งไว้ล่วงหน้า เพราะบูตครั้งแรกใช้เวลานาน และเตรียมข้อมูลตัวอย่างที่มี Streak หลายวันไว้ก่อน\n"
      "•  ภาพหลักฐานการทดสอบทุกขั้นอยู่ในโฟลเดอร์ docs/test_evidence เปิดให้กรรมการดูได้\n"
      "•  สมาชิกทุกคนต้องรู้ภาพรวมโครงงาน และต้องอธิบายงานที่ตนเองรับผิดชอบได้\n"
      "•  งานที่ยังไม่เสร็จให้ระบุอย่างตรงไปตรงมา พร้อมแผนดำเนินการต่อในภาคเรียนหน้าเพื่อเข้าสู่การสอบจบ",
      14.5, False, DARK, space_after=13, line=1.15)
rect(s, 1.1, 5.4, 11.13, 1.15, RGBColor(0xFD, 0xEF, 0xEF), RGBColor(0xE5, 0xB4, 0xB4), adj=0.15)
txbox(s, 1.5, 5.65, 10.3, 0.7,
      "หลักสำคัญ : “คะแนนโครงการ 1 พิจารณาจากสิ่งที่พัฒนาได้จริง ความเข้าใจของผู้จัดทำ คุณภาพหลักฐาน\n"
      "และความพร้อมที่จะพัฒนาต่อให้สมบูรณ์ในโครงการ 2”",
      14, True, RGBColor(0xB0, 0x3A, 0x3A), align=PP_ALIGN.CENTER, line=1.25)

out_dir = os.path.join("D:\\New folder\\habit_tracker", "docs")
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, "นำเสนอโครงการ1_HabitTracker.pptx")
prs.save(out)
print("SAVED:", out)
print("size:", os.path.getsize(out), "bytes | slides:", len(prs.slides._sldIdLst))
