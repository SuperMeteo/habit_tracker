# -*- coding: utf-8 -*-
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "Noto Looped Thai"
NAVY = RGBColor(0x0C, 0x2D, 0x48)
NAVY2 = RGBColor(0x17, 0x4B, 0x70)
NAVY3 = RGBColor(0x10, 0x3A, 0x59)
GREY = RGBColor(0x64, 0x74, 0x8B)
INK = RGBColor(0x1F, 0x29, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xDD, 0xE3, 0xEA)
TEAL = RGBColor(0x2E, 0x8B, 0x8B)
TEALBG = RGBColor(0xE9, 0xF5, 0xF5)
TEALLN = RGBColor(0xBF, 0xE0, 0xE0)
GREEN = RGBColor(0x2F, 0x85, 0x5A)
GREENBG = RGBColor(0xE7, 0xF5, 0xED)
RED = RGBColor(0xC5, 0x30, 0x30)
AMBER = RGBColor(0xD6, 0x9E, 0x2E)
AMBERBG = RGBColor(0xFF, 0xF8, 0xE6)
AMBERLN = RGBColor(0xF3, 0xD3, 0x8A)
AMBERTX = RGBColor(0x8A, 0x5B, 0x00)
PALE = RGBColor(0xD6, 0xE4, 0xEE)
ICE = RGBColor(0xBC, 0xE5, 0xE5)


def blank_from(template, out_unused=None):
    prs = Presentation(template)
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    return prs


def txt(slide, x, y, w, h, text, size=12, bold=False, color=INK,
        align=PP_ALIGN.LEFT, line=1.15, space=4, anchor=MSO_ANCHOR.TOP, font=FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = ln
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(0.75)
    if adj is not None and len(s.adjustments):
        try: s.adjustments[0] = adj
        except Exception: pass
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def head(slide, num, title, subtitle, key):
    box(slide, 0.55, 0.42, 0.6, 0.5, NAVY, NAVY, adj=0.15)
    txt(slide, 0.55, 0.49, 0.6, 0.24, num, 13, True, WHITE, PP_ALIGN.CENTER)
    txt(slide, 1.3, 0.36, 10.8, 0.42, title, 25, True, NAVY)
    txt(slide, 1.3, 0.81, 10.7, 0.32, subtitle, 12.5, False, GREY)
    box(slide, 9.75, 0.42, 3.03, 0.67, TEALBG, TEALLN, adj=0.18)
    txt(slide, 9.95, 0.49, 2.65, 0.2, "ประเด็นสำคัญ", 10, True, TEAL)
    txt(slide, 9.95, 0.7, 2.65, 0.28, key, 9.5, False, INK)


def card(slide, x, y, w, h, bar, title, body, body_size=12, title_size=16, num=None):
    box(slide, x, y, w, h, WHITE, BORDER, adj=0.05)
    box(slide, x, y, 0.06, h, bar, bar, shape=MSO_SHAPE.RECTANGLE, adj=None)
    tx = x + 0.25
    if num is not None:
        box(slide, x + 0.25, y + 0.22, 0.46, 0.46, bar, bar, shape=MSO_SHAPE.OVAL, adj=None)
        txt(slide, x + 0.25, y + 0.37, 0.46, 0.18, num, 11.5, True, WHITE, PP_ALIGN.CENTER)
        tx = x + 0.87
    txt(slide, tx, y + 0.24, w - (tx - x) - 0.2, 0.34, title, title_size, True, NAVY)
    if body:
        txt(slide, x + 0.25, y + 0.78, w - 0.5, h - 0.95, body, body_size, False, INK, line=1.25, space=3)


def note(slide, x, y, w, h, label, body, body_size=13.2):
    box(slide, x, y, w, h, AMBERBG, AMBERLN, adj=0.12)
    txt(slide, x + 0.28, y + 0.29, 1.7, 0.28, label, 15, True, AMBERTX)
    txt(slide, x + 1.85, y + 0.17, w - 2.2, h - 0.35, body, body_size, False, INK, line=1.35, space=4)


def table(slide, x, y, w, headers, rows, col_w, row_h=0.42, head_h=0.44,
          head_size=11.5, body_size=11):
    shp = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                                 Inches(w), Inches(head_h + row_h * len(rows)))
    t = shp.table
    t.first_row = True
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    t.rows[0].height = Inches(head_h)
    for i in range(1, len(rows) + 1):
        t.rows[i].height = Inches(row_h)

    def fill(cell, text, size, bold, color, bg):
        cell.fill.solid(); cell.fill.fore_color.rgb = bg
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.line_spacing = 1.1
        r = p.add_run(); r.text = text
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color

    for c, h in enumerate(headers):
        fill(t.cell(0, c), h, head_size, True, WHITE, NAVY)
    for ri, row in enumerate(rows, start=1):
        bg = WHITE if ri % 2 else RGBColor(0xF4, 0xF7, 0xFA)
        for ci, val in enumerate(row):
            text, color, bold = val if isinstance(val, tuple) else (val, INK, False)
            fill(t.cell(ri, ci), text, body_size, bold, color, bg)
    return t


def picture(slide, path, x, y, h):
    from PIL import Image
    with Image.open(path) as im:
        w = h * im.width / im.height
    box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, WHITE, BORDER, adj=0.05)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def picture_w(slide, path, x, y, w):
    from PIL import Image
    with Image.open(path) as im:
        h = w * im.height / im.width
    box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, WHITE, BORDER, adj=0.05)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return h
